"""Contract tests for `MRNAStabilitySource` (T102) on synthetic Agarwal/Wu/TargetScan files.

The parsers are pinned on the properties that decide the numbers:

* Table S2's header is FOUND, not assumed -- row 1 of the `human` sheet is a note and row 2
  is the header, and a reader that trusted that would eat the first gene if the note ever
  went away;
* the 54 sample columns are grouped into cell lines by an exact token first and a substring
  only as a fallback, a column matching neither is left UNASSIGNED rather than guessed at,
  and the whole grouping is written into LINEAGE.json;
* the per-line `_n` is the per-GENE count of samples the mean was taken over, not the number
  of columns in the group;
* an ORF is ungapped before it is read as codons (`-` and `.` both), and the mean CSC is over
  the 61 SENSE codons weighted by how often each occurs -- hand-computed here for a 3-codon
  ORF;
* the frame checks (multiple of 3, starts ATG, internal stops) are COUNTED and reported, and
  no gene is dropped for failing one;
* the four Ensembl-keyed decay sheets are read raw, their medians differ in sign, and that
  difference is recorded rather than standardised away;
* every table leaves a LINEAGE entry, and two tables sharing a derived directory merge
  rather than overwrite;
* `build()` is a node_feature aligned to the gene space with a `has_value` mask, because a
  CSC of 0.0 is a measurement and a gene with no row is also 0.0.
"""
from __future__ import annotations

import json
import zipfile

import numpy as np
import openpyxl
import pandas as pd
import pytest

from sidechain.priors.posttx_stability import (
    CODON_COLUMNS,
    CSC_COLUMNS,
    MRNAStabilitySource,
    codon_features,
    group_sample_columns,
    ungap_orf,
)

# ---------------------------------------------------------------- the CSC table --

# Every codon scores 0.0 unless it is named here, so a gene's mean is the mean over the
# few codons that matter and can be checked by hand.
CSC_VALUES: dict[str, dict[str, float]] = {
    "ATG": {"293T_endo": 0.30, "HeLa_endo": 0.20, "RPE_endo": 0.10,
            "293T_ORFome": -0.10, "K562_ORFome": 0.05, "K562_SLAM": 0.00},
    "GCT": {"293T_endo": 0.60, "HeLa_endo": 0.30, "RPE_endo": 0.30,
            "293T_ORFome": 0.40, "K562_ORFome": 0.20, "K562_SLAM": 0.10},
    "AAA": {"293T_endo": -0.30, "HeLa_endo": -0.20, "RPE_endo": -0.10,
            "293T_ORFome": -0.20, "K562_ORFome": -0.15, "K562_SLAM": -0.05},
}
STOPS = ("TAA", "TAG", "TGA")
ALL_CODONS = [a + b + c for a in "ACGT" for b in "ACGT" for c in "ACGT"]


def _csc_csv() -> str:
    # the real file's first column header is the one field nobody has measured, so the
    # fixture gives it an unhelpful name: the reader has to find the codons by content
    head = "X," + ",".join(CSC_COLUMNS) + ",genome_count,transcriptome_count\n"
    rows = []
    for codon in ALL_CODONS:
        if codon in STOPS:
            continue                                    # 61 sense codons, as published
        values = CSC_VALUES.get(codon, {})
        cells = [f"{values.get(c, 0.0):.4f}" for c in CSC_COLUMNS]
        rows.append(f"{codon}," + ",".join(cells) + ",1000,500\n")
    return head + "".join(rows)


# ------------------------------------------------------------- TargetScan inputs --

GENE_INFO = (
    "Transcript ID\tGene ID\tGene symbol\tGene description\tSpecies ID\t3P-seq tags\tRepresentative transcript?\n"
    "ENST1.3\tENSG1.7\tA1BG\tdesc\t9606\t74\t1\n"
    "ENST2.2\tENSG2.10\tA1CF\tdesc\t9606\t80\t1\n"
    "ENST3.1\tENSG2.10\tA1CF\tdesc\t9606\t2\t0\n"          # not representative: ignored
    "ENST4.1\tENSG4.1\tNOSTART\tdesc\t9606\t5\t1\n"
    "ENST5.1\tENSG5.1\tFRAME\tdesc\t9606\t5\t1\n"
    "ENSMUST1.1\tENSMUSG1.1\tA1bg\tdesc\t10090\t1\t1\n"    # mouse: ignored
)

# ENST1: the hand-computed ORF. Ungapped ATG GCT AAA TAA -- three sense codons and a
# terminal stop, with alignment gaps in the middle of the start codon.
# ENST2: a `.` gap and an INTERNAL stop -- ATG AAA TGA AAA TAA.
# ENST4: does not start ATG.  ENST5: eight nucleotides, so not a whole number of codons.
ORF_HEADER = "Transcript ID\tGene ID\tGene Symbol\tSpecies ID\tORF Sequence\n"
ORF_ROWS = (
    "ENST1.3\tENSG1.7\tA1BG\t9606\tAT--GGCTAAATAA\n"
    "ENST1.3\tENSG1.7\tA1BG\t10090\tATGAAAAAATAA\n"        # mouse row of the same site: ignored
    "ENST2.2\tENSG2.10\tA1CF\t9606\tATG..AAATGAAAATAA\n"
    "ENST3.1\tENSG2.10\tA1CF\t9606\tATGGGGTAA\n"           # not representative: ignored
    "ENST4.1\tENSG4.1\tNOSTART\t9606\tGGGAAATAA\n"
    "ENST5.1\tENSG5.1\tFRAME\t9606\tATGGCTAA\n"
)
ORF_SEQUENCES = ORF_HEADER + ORF_ROWS
# the same rows under TargetScan's other spelling and a different column order: the reader
# resolves columns by name, so this must parse to exactly the same table
ORF_SEQUENCES_ALT = (
    "Gene ID\tRefseq ID\tSpecies ID\tGene Symbol\tORF sequence\n"
    + "".join("\t".join([p[1], p[0], p[3], p[2], p[4]]) + "\n"
              for p in (line.split("\t") for line in ORF_ROWS.strip("\n").split("\n")))
)

# the layout the real release uses (read on the box 2026-09-28): no header, three columns --
# versioned transcript id, species id, gapped lower-case ORF; gene id and symbol come from Gene_info
ORF_SEQUENCES_HEADERLESS = "".join(
    "\t".join([p[0], p[3], p[4].lower()]) + "\n"
    for p in (line.split("\t") for line in ORF_ROWS.strip("\n").split("\n")))

# ENST1, by hand: sense codons ATG, GCT, AAA, one of each.
CSC_ENST1 = {c: (CSC_VALUES["ATG"][c] + CSC_VALUES["GCT"][c] + CSC_VALUES["AAA"][c]) / 3
             for c in CSC_COLUMNS}
ENDO_ENST1 = sum(CSC_ENST1[c] for c in ("293T_endo", "HeLa_endo", "RPE_endo")) / 3
# ENST2: ATG AAA TGA | AAA TAA -- the CDS stops at the first in-frame TGA, so ATG and AAA.
CSC_ENST2 = {c: (CSC_VALUES["ATG"][c] + CSC_VALUES["AAA"][c]) / 2 for c in CSC_COLUMNS}
ENDO_ENST2 = sum(CSC_ENST2[c] for c in ("293T_endo", "HeLa_endo", "RPE_endo")) / 3

# --------------------------------------------------------------- the half-life xlsx --

HALFLIFE_NOTE = ("Note: The half-life (PC1) is a summary of all of the datasets combined, "
                 "excluding data from Gejman et al.")
HALFLIFE_HEADER = ["Ensembl Gene Id", "Gene name", "half-life (PC1)",
                   "Bazzini_ActD_HEK293_1", "Bazzini_ActD_HEK293T_2",
                   "Schofield_SLAMseq_K562_1", "Wu_ActD_HeLa_1", "Gejman_LCL_1",
                   "Mystery_XYZ_1"]
HALFLIFE_ROWS = [
    ["ENSG1.7", "A1BG", 3.5, 1.0, 3.0, 5.0, 7.0, 9.0, 11.0],
    ["ENSG2.10", "A1CF", 4.5, 4.0, None, 6.0, 8.0, 10.0, 12.0],   # one HEK293 sample missing
    ["ENSG4.1", "NOSTART", None, 2.0, 2.0, 2.0, 2.0, 2.0, 2.0],   # no consensus value
    ["ENSG9.1", "ONLYHL", 2.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0],
    ["ENSG1.7", "A1BG", 99.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0],      # duplicate gene id
    ["NOTANID", "JUNK", 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0],       # not an ENSG: dropped
]

# ------------------------------------------------------------------- the decay xlsx --

DECAY_SHEET_ROWS = {
    "293T-endogenous": [("ENSG1.7", 0.010), ("ENSG2.10", 0.014), ("ENSG4.1", 0.012)],
    "HeLa-endogenous": [("ENSG1.7", 0.020), ("ENSG2.10", 0.030)],
    "RPE-endogenous": [("ENSG1.7", 0.011), ("ENSG2.10", 0.013)],
    "k562-SLAM-seq": [("ENSG1.7", -0.10), ("ENSG2.10", -0.12), ("ENSG4.1", -0.14)],
}
DECAY_ORFOME_ROWS = [("ACGTACGTACGTACGTACGTACGT", 0.5), ("TTTTTTTTTTTTTTTTTTTTTTTT", 0.6)]


# ------------------------------------------------------------------------ fixtures --

def _zip(path, member, text):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(member, text)


def _provenance(dest, names):
    (dest / "PROVENANCE.json").write_text(json.dumps({"selected": [
        {"name": n, "size_bytes": (dest / n).stat().st_size, "checksum": "http-last-modified:x"}
        for n in names]}))


def _write_halflife_xlsx(path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "human"
    sheet.append([HALFLIFE_NOTE])
    sheet.append(HALFLIFE_HEADER)
    for row in HALFLIFE_ROWS:
        sheet.append(row)
    mouse = book.create_sheet("mouse")
    mouse.append(["Note: the mouse half-lives"])
    mouse.append(["Ensembl Gene Id", "Gene name", "half-life (PC1)", "Mouse_ActD_ESC_1"])
    mouse.append(["ENSMUSG1.1", "A1bg", 1.0, 1.0])
    book.save(path)


def _write_decay_xlsx(path):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in DECAY_SHEET_ROWS.items():
        sheet = book.create_sheet(name)
        sheet.append(["gene_id", "decay_rate", "r_squared"])
        for gene, rate in rows:
            sheet.append([gene, rate, 0.9])
    for name in ("293T-ORFome", "k562-ORFome"):
        sheet = book.create_sheet(name)
        sheet.append(["barcode", "decay_rate"])
        for barcode, rate in DECAY_ORFOME_ROWS:
            sheet.append([barcode, rate])
    book.save(path)


@pytest.fixture
def sources(tmp_path):
    """(halflife source, codon source) over a synthetic release, sharing one test registry."""
    root = tmp_path
    springer = root / "external" / "springer"
    elife = root / "external" / "elife"
    targetscan = root / "external" / "ts"
    for d in (springer, elife, targetscan):
        d.mkdir(parents=True)

    _write_halflife_xlsx(springer / "13059_2022_2811_MOESM3_ESM.xlsx")
    _provenance(springer, ["13059_2022_2811_MOESM3_ESM.xlsx"])

    (elife / "elife-45396-fig1-data2-v2.csv").write_text(_csc_csv())
    _write_decay_xlsx(elife / "elife-45396-fig1-data1-v2.xlsx")
    _provenance(elife, ["elife-45396-fig1-data2-v2.csv", "elife-45396-fig1-data1-v2.xlsx"])

    _zip(targetscan / "Gene_info.txt.zip", "Gene_info.txt", GENE_INFO)
    _zip(targetscan / "ORF_Sequences.txt.zip", "ORF_Sequences.txt", ORF_SEQUENCES)
    _provenance(targetscan, ["Gene_info.txt.zip", "ORF_Sequences.txt.zip"])

    registry = root / "registry.yaml"
    registry.write_text(
        "master_gene_space:\n  reference: gencode\n  version: \"v47\"\n"
        "  id_type: ensembl_gene_id\n"
        "sources:\n"
        "  - name: halflife_consensus\n    kind: node_feature\n    layer: posttx\n"
        "    relation: mrna_halflife\n    loader: MRNAStabilitySource\n    enabled: false\n"
        "    dest: external/springer\n    derived: derived/halflife\n    license: CC-BY-4.0\n"
        "  - name: codon_optimality_wu2019\n    kind: node_feature\n    layer: posttx\n"
        "    relation: codon_optimality\n    loader: MRNAStabilitySource\n    enabled: false\n"
        "    dest: external/elife\n    derived: derived/codon\n    license: CC-BY-4.0\n"
        "  - name: targetscan\n    kind: edge\n    layer: posttx\n    relation: mirna_target\n"
        "    loader: MiRNATargetSource\n    enabled: false\n"
        "    dest: external/ts\n    derived: derived/ts\n"
        "    license: Free-for-research-with-citation\n"
    )

    def build(name):
        import yaml
        spec = next(s for s in yaml.safe_load(registry.read_text())["sources"]
                    if s["name"] == name)
        return MRNAStabilitySource(spec, {}, root=root, registry=registry)

    return build("halflife_consensus"), build("codon_optimality_wu2019")


# ------------------------------------------------------------------- the half-life --

def test_halflife_header_is_found_below_the_note_row(sources):
    halflife, _ = sources
    table = halflife.halflife_table(rebuild=True)
    assert list(table.columns)[:3] == ["gene_id", "symbol", "halflife_pc1"]
    assert set(table["gene_id"]) == {"ENSG1", "ENSG2", "ENSG4", "ENSG9"}   # NOTANID dropped
    lineage = json.loads((halflife.derived / "LINEAGE.json").read_text())
    entry = lineage["entries"]["halflife-agarwal-kelley-2022/halflife"]
    assert entry["counts"]["header_row_index"] == 1
    assert entry["header_rows_above"].startswith("Note: The half-life (PC1)")


def test_halflife_values_lines_and_per_gene_counts(sources):
    halflife, _ = sources
    table = halflife.halflife_table(rebuild=True).set_index("gene_id")
    assert table.loc["ENSG1", "halflife_pc1"] == pytest.approx(3.5)        # as published, not z-scored
    assert table.loc["ENSG1", "halflife_hek293_mean"] == pytest.approx(2.0)   # (1.0 + 3.0) / 2
    assert table.loc["ENSG1", "halflife_hek293_n"] == 2
    # ENSG2 is missing one of its two HEK293 samples: the mean is over the one it has
    assert table.loc["ENSG2", "halflife_hek293_mean"] == pytest.approx(4.0)
    assert table.loc["ENSG2", "halflife_hek293_n"] == 1
    assert table.loc["ENSG1", "halflife_k562_mean"] == pytest.approx(5.0)
    assert table.loc["ENSG1", "halflife_hela_mean"] == pytest.approx(7.0)
    assert table.loc["ENSG1", "halflife_lcl_mean"] == pytest.approx(9.0)
    assert np.isnan(table.loc["ENSG4", "halflife_pc1"])
    # the duplicate ENSG1 row (pc1 99.0) is dropped, not averaged in
    assert table.loc["ENSG1", "halflife_pc1"] != pytest.approx(99.0)


def test_halflife_lineage_records_the_grouping_and_the_unassigned_column(sources):
    halflife, _ = sources
    halflife.halflife_table(rebuild=True)
    entry = json.loads((halflife.derived / "LINEAGE.json").read_text())[
        "entries"]["halflife-agarwal-kelley-2022/halflife"]
    grouping = entry["sample_column_grouping"]
    assert grouping["hek293"] == ["Bazzini_ActD_HEK293T_2", "Bazzini_ActD_HEK293_1"]
    assert grouping["k562"] == ["Schofield_SLAMseq_K562_1"]
    assert grouping["hela"] == ["Wu_ActD_HeLa_1"] and grouping["lcl"] == ["Gejman_LCL_1"]
    assert entry["sample_columns_unassigned"] == ["Mystery_XYZ_1"]
    counts = entry["counts"]
    assert counts["sample_columns"] == 6 and counts["sample_columns_unassigned"] == 1
    assert counts["genes_out"] == 4 and counts["genes_with_pc1"] == 3
    assert counts["duplicate_gene_ids_dropped"] == 1
    assert set(entry["sample_column_match_rule"].values()) == {"token"}


def test_group_sample_columns_falls_back_to_substrings_and_refuses_to_guess():
    by_line, how, unassigned = group_sample_columns(
        ["LabHEK293Tsample", "run_K562_3", "Smith_H1975_1", "plate1"])
    assert by_line["hek293"] == ["LabHEK293Tsample"] and how["LabHEK293Tsample"] == "substring"
    assert by_line["k562"] == ["run_K562_3"] and how["run_K562_3"] == "token"
    # H1975 is a lung line, not H1: the short `H1` substring is deliberately not a rule
    assert unassigned == ["Smith_H1975_1", "plate1"]


# ------------------------------------------------------------------ the codon score --

def test_ungap_orf_removes_every_gap_symbol_and_normalises_the_alphabet():
    assert ungap_orf("AT--GGC.TAA") == "ATGGCTAA"
    assert ungap_orf("aug~gcu\tuaa") == "ATGGCTTAA"


def test_codon_features_of_a_three_codon_orf_by_hand(sources):
    _, codon = sources
    csc, resolved = codon.csc_table()
    assert len(csc) == 61 and list(csc.columns) == list(CSC_COLUMNS)
    assert resolved["293T_endo"] == "293T_endo"
    features = codon_features("ATGGCTAAATAA", csc)
    assert features["orf_len"] == 12 and features["n_codons"] == 4
    assert features["n_sense_codons"] == 3 and features["n_stops_raw"] == 1
    assert features["has_in_frame_stop"] and features["len_multiple_of_3"]
    assert features["trailing_nt"] == 0
    assert features["starts_atg"] and features["n_ambiguous_codons"] == 0
    assert features["gc3"] == pytest.approx(1 / 3)          # ATG ends G; GCT and AAA do not
    assert features["codon_ATG"] == 1 and features["codon_TAA"] == 1
    for column in CSC_COLUMNS:
        assert features[f"csc_{column}"] == pytest.approx(CSC_ENST1[column])
    assert features["csc_endo_mean"] == pytest.approx(ENDO_ENST1)


def test_codon_table_keeps_human_representative_transcripts_and_scores_them(sources):
    _, codon = sources
    raw = codon.codon_table(rebuild=True)
    assert list(raw.columns)[:len(CODON_COLUMNS)] == list(CODON_COLUMNS)
    table = raw.set_index("gene_id")
    # ENST3 is not representative, the mouse row of ENST1 is not human
    assert set(table.index) == {"ENSG1", "ENSG2", "ENSG4", "ENSG5"}
    assert table.loc["ENSG1", "orf_len"] == 12          # the two `-` gaps are removed first
    assert table.loc["ENSG1", "csc_endo_mean"] == pytest.approx(ENDO_ENST1)
    # `.` is a gap too; ATG AAA TGA | AAA TAA -- the CDS ends at the FIRST in-frame stop
    assert table.loc["ENSG2", "orf_len"] == 9 and table.loc["ENSG2", "raw_len"] == 15
    assert table.loc["ENSG2", "trailing_nt"] == 6 and table.loc["ENSG2", "n_stops_raw"] == 2
    assert table.loc["ENSG2", "csc_endo_mean"] == pytest.approx(ENDO_ENST2)


def test_codon_table_counts_the_validity_checks_without_dropping_a_gene(sources):
    _, codon = sources
    table = codon.codon_table(rebuild=True).set_index("gene_id")
    assert not bool(table.loc["ENSG4", "starts_atg"])          # GGG AAA TAA
    assert not bool(table.loc["ENSG5", "len_multiple_of_3"])   # eight nucleotides
    assert not bool(table.loc["ENSG5", "has_in_frame_stop"])
    assert np.isfinite(table.loc["ENSG4", "csc_endo_mean"])    # scored anyway
    counts = json.loads((codon.derived / "LINEAGE.json").read_text())[
        "entries"]["codon-optimality-wu2019/codon_score"]["counts"]
    assert counts["orf_rows_human"] == 5 and counts["representative_transcripts"] == 4
    assert counts["genes_out"] == 4 and counts["codons_scored"] == 61
    assert counts["genes_raw_length_multiple_of_3"] == 3 and counts["genes_starting_atg"] == 3
    assert counts["genes_with_in_frame_stop"] == 3 and counts["genes_trailing_over_60nt"] == 0


def test_codon_table_reads_the_other_orf_header_spelling(sources, tmp_path):
    _, codon = sources
    _zip(tmp_path / "external" / "ts" / "ORF_Sequences.txt.zip", "ORF_Sequences.txt",
         ORF_SEQUENCES_ALT)
    table = codon.codon_table(rebuild=True).set_index("gene_id")
    assert set(table.index) == {"ENSG1", "ENSG2", "ENSG4", "ENSG5"}
    assert table.loc["ENSG1", "csc_endo_mean"] == pytest.approx(ENDO_ENST1)


def test_codon_features_stop_at_the_first_in_frame_stop_like_the_real_release(sources):
    """ARF5's shape: the row runs 14 nt past the CDS's TAA; those bases are not codons."""
    _, codon = sources
    csc, _ = codon.csc_table()
    features = codon_features("ATGGCTAAATAA" + "CCAGCCAGGGGCAG", csc)
    assert features["orf_len"] == 12 and features["trailing_nt"] == 14
    assert features["n_sense_codons"] == 3
    assert features["csc_endo_mean"] == pytest.approx(ENDO_ENST1)


def test_codon_table_reads_the_real_headerless_release(sources, tmp_path):
    _, codon = sources
    _zip(tmp_path / "external" / "ts" / "ORF_Sequences.txt.zip", "ORF_Sequences.txt",
         ORF_SEQUENCES_HEADERLESS)
    table = codon.codon_table(rebuild=True).set_index("gene_id")
    assert set(table.index) == {"ENSG1", "ENSG2", "ENSG4", "ENSG5"}
    assert table.loc["ENSG1", "symbol"] == "A1BG"
    assert table.loc["ENSG1", "orf_len"] == 12
    assert table.loc["ENSG1", "csc_endo_mean"] == pytest.approx(ENDO_ENST1)


def test_codon_lineage_names_both_publishers_inputs(sources):
    _, codon = sources
    codon.codon_table(rebuild=True)
    entry = json.loads((codon.derived / "LINEAGE.json").read_text())[
        "entries"]["codon-optimality-wu2019/codon_score"]
    inputs = entry["inputs"]
    assert set(inputs) == {"elife-45396-fig1-data2-v2.csv", "ORF_Sequences.txt.zip",
                           "Gene_info.txt.zip"}
    assert inputs["ORF_Sequences.txt.zip"]["dest"].endswith("external/ts")
    assert inputs["elife-45396-fig1-data2-v2.csv"]["dest"].endswith("external/elife")
    assert all(len(v["sha256"]) == 64 for v in inputs.values())


# ------------------------------------------------------------------ the decay rates --

def test_decay_table_reads_the_ensembl_sheets_and_skips_the_barcode_ones(sources):
    _, codon = sources
    table = codon.decay_table(rebuild=True).set_index("gene_id")
    assert set(table.columns) == {"decay_rate_293T_endogenous", "decay_rate_HeLa_endogenous",
                                  "decay_rate_RPE_endogenous", "decay_rate_k562_SLAM_seq"}
    assert set(table.index) == {"ENSG1", "ENSG2", "ENSG4"}
    assert table.loc["ENSG1", "decay_rate_293T_endogenous"] == pytest.approx(0.010)
    assert table.loc["ENSG1", "decay_rate_k562_SLAM_seq"] == pytest.approx(-0.10)
    # ENSG4 has no HeLa row: outer join leaves NaN rather than dropping the gene
    assert np.isnan(table.loc["ENSG4", "decay_rate_HeLa_endogenous"])


def test_decay_lineage_records_the_opposite_signs_rather_than_standardising(sources):
    _, codon = sources
    codon.decay_table(rebuild=True)
    entry = json.loads((codon.derived / "LINEAGE.json").read_text())[
        "entries"]["codon-optimality-wu2019/decay_rates"]
    sheets = entry["sheets"]
    assert sheets["293T-endogenous"]["median"] == pytest.approx(0.012)
    assert sheets["293T-endogenous"]["median_sign"] == "positive"
    assert sheets["k562-SLAM-seq"]["median"] == pytest.approx(-0.12)
    assert sheets["k562-SLAM-seq"]["median_sign"] == "negative"
    assert sheets["k562-SLAM-seq"]["line"] == "K562"        # the Methods name, not the ATCC number
    assert set(entry["sheets_skipped"]) == {"293T-ORFome", "k562-ORFome"}
    assert "does not" in entry["sign_convention"]


def test_two_tables_in_one_derived_directory_merge_their_lineage_entries(sources):
    _, codon = sources
    codon.codon_table(rebuild=True)
    codon.decay_table(rebuild=True)
    entries = json.loads((codon.derived / "LINEAGE.json").read_text())["entries"]
    assert {"codon-optimality-wu2019/codon_score",
            "codon-optimality-wu2019/decay_rates"} <= set(entries)


def test_tables_are_cached_and_rebuild_is_opt_in(sources):
    halflife, codon = sources
    first = halflife.halflife_table(rebuild=True)
    assert (halflife.derived / "halflife.parquet").exists()
    assert len(halflife.halflife_table()) == len(first)
    codon.codon_table(rebuild=True)
    assert (codon.derived / "codon_score.parquet").exists()


# ------------------------------------------------------------------------- build() --

def _with_index(source, gene_index):
    return MRNAStabilitySource(source.spec, gene_index, root=source.root,
                               registry=source.registry)


def test_build_is_a_node_feature_the_contract_accepts_with_a_value_mask(sources):
    halflife, _ = sources
    source = _with_index(halflife, {"ENSG1": 0, "ENSG2": 1, "ENSG4": 2, "ENSG9": 3})
    art = source.build_checked()
    assert art.kind == "node_feature" and art.layer == "posttx"
    assert art.features.shape == (4, 2) and art.features.dtype == np.float32
    assert art.meta["feature_names"] == ["halflife_pc1", "csc_endo_mean"]

    mask = art.meta["has_value"]
    assert mask.shape == (4, 2) and mask.dtype == bool
    assert art.features[0, 0] == pytest.approx(3.5) and mask[0, 0]
    assert art.features[1, 1] == pytest.approx(ENDO_ENST2, abs=1e-6) and mask[1, 1]
    # ENSG4 has no consensus half-life: a zero the mask says is NOT a measurement
    assert art.features[2, 0] == 0.0 and not mask[2, 0]
    # ENSG9 has a half-life but no ORF, so no codon score
    assert mask[3, 0] and not mask[3, 1] and art.features[3, 1] == 0.0
    assert art.meta["genes_with_value"] == {"halflife_pc1": 3, "csc_endo_mean": 3}
    assert art.meta["genes_with_any_value"] == 4


def test_build_works_from_either_block_and_drops_genes_off_the_axis(sources):
    halflife, codon = sources
    from_codon = _with_index(codon, {"ENSG1": 0}).build_checked()
    from_halflife = _with_index(halflife, {"ENSG1": 0}).build_checked()
    assert from_codon.features.shape == (1, 2)
    np.testing.assert_allclose(from_codon.features, from_halflife.features)
    # the artifact carries the relation of the block it was built from; both are named
    assert from_codon.relation == "codon_optimality"
    assert from_halflife.relation == "mrna_halflife"
    assert from_codon.meta["relations"]["halflife_pc1"] == "mrna_halflife"
    # ENSG5 and the rest are off the axis: 4 candidate codon rows, 1 kept
    assert from_codon.meta["candidate_genes"]["csc_endo_mean"] == 4
    assert from_codon.meta["genes_with_value"]["csc_endo_mean"] == 1


def test_registry_resolves_the_loader_for_both_blocks():
    from sidechain.data.registry import LOADERS
    from sidechain.priors.posttx_mirna import spec_from_registry

    assert LOADERS["MRNAStabilitySource"] is MRNAStabilitySource
    for name in ("halflife_consensus", "codon_optimality_wu2019"):
        spec = spec_from_registry(name)
        assert spec["loader"] == "MRNAStabilitySource" and spec["kind"] == "node_feature"
        assert spec["enabled"] is False and spec["license"] == "CC-BY-4.0"
        assert spec["allow_missing_checksum"] is True
        assert "license_override_source" in spec
        source = LOADERS[spec["loader"]](spec=spec, gene_index={})
        assert source.layer == "posttx"


def test_missing_csc_column_raises_rather_than_scoring_on_what_is_left(sources, tmp_path):
    _, codon = sources
    text = _csc_csv().replace("K562_SLAM", "K562_SOMETHING_ELSE")
    (tmp_path / "external" / "elife" / "elife-45396-fig1-data2-v2.csv").write_text(text)
    with pytest.raises(ValueError, match="missing CSC column"):
        codon.csc_table()


def test_a_header_without_the_key_column_raises(sources, tmp_path):
    halflife, _ = sources
    path = tmp_path / "external" / "springer" / "13059_2022_2811_MOESM3_ESM.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "human"
    sheet.append(["Note"])
    sheet.append(["gene", "name", "half-life (PC1)"])
    sheet.append(["ENSG1", "A1BG", 1.0])
    book.save(path)
    with pytest.raises(ValueError, match="ensembl gene id"):
        halflife.halflife_table(rebuild=True)


def test_missing_decay_sheet_raises_with_the_sheets_it_found(sources, tmp_path):
    _, codon = sources
    path = tmp_path / "external" / "elife" / "elife-45396-fig1-data1-v2.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "293T-endogenous"
    book.active.append(["gene_id", "decay_rate"])
    book.active.append(["ENSG1", 0.01])
    book.save(path)
    with pytest.raises(ValueError, match="missing sheet"):
        codon.decay_table(rebuild=True)


def test_pandas_frames_round_trip_through_parquet(sources):
    """The cached parquet is the table, not a lossy copy of it."""
    _, codon = sources
    built = codon.codon_table(rebuild=True)
    cached = pd.read_parquet(codon.derived / "codon_score.parquet")
    pd.testing.assert_frame_equal(built, cached)
